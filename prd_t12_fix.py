# -*- coding: utf-8 -*-
"""T-12：从《需求规格说明书 .docx》删除昼夜机制表述（代码不存在昼夜逻辑，文档向代码靠）。

只删昼夜相关分句，保留段落/单元格其余内容。运行前备份 PRD（带时间戳）。
"""
import shutil
import time

import docx

PRD_PATH = r'C:/Users/jeezeqi/Desktop/需求规格说明书 .docx'
BAK_PATH = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/_backup/需求规格说明书_{}.bak.docx'.format(
    time.strftime('%Y%m%d_%H%M%S'))

# (旧子串, 新子串) —— 段落与表格单元格通用
REPLACEMENTS = [
    ('推进昼夜：当前为白天则切到黑夜；当前为黑夜则切到白天并完成一次完整结算。',
     '每个游戏日完成一次完整结算'),
    ('仅在白天由玩家在建筑详情弹窗点击', '由玩家在建筑详情弹窗点击'),
    ('修复成功直接回到运行状态；修复操作在黑夜禁用。', '修复成功直接回到运行状态。'),
    ('校验建筑容量、空闲人口、当前是否处于白天', '校验建筑容量、空闲人口'),
    ('分配给不支持居住建筑、黑夜期间操作：拒绝操作', '分配给不支持居住建筑：拒绝操作'),
    ('每个游戏日触发一次昼夜更替以及每日结算', '每个游戏日触发一次完整结算'),
    ('资源不足恢复阻断、黑夜禁止修复：操作拒绝', '资源不足恢复阻断：操作拒绝'),
    ('资源数量；昼夜状态指示（白天 / 黑夜）；所有错误提示弹窗', '资源数量；所有错误提示弹窗'),
    ('（含剩余天数、昼夜限制）', '（含剩余天数）'),
]


def replace_in_par(p, old, new):
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    p.runs[0].text = p.text.replace(old, new)
    for r in p.runs[1:]:
        r.text = ''
    return True


def main():
    shutil.copyfile(PRD_PATH, BAK_PATH)
    print('已备份:', BAK_PATH)

    d = docx.Document(PRD_PATH)

    # 1) 通用子串替换（段落 + 表格单元格）
    for old, new in REPLACEMENTS:
        n = 0
        for p in d.paragraphs:
            if replace_in_par(p, old, new):
                n += 1
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        if replace_in_par(p, old, new):
                            n += 1
        print(f'替换 [{old[:24]}...] -> {n} 处')
        assert n >= 1, '未命中: ' + old

    # 2) 三个需要整句删除的单元格（删掉分句，保留其余内容）
    special = [
        ('黑夜期间不允许放置新建筑', '；黑夜期间'),
        ('黑夜期间分配 / 移除按钮置灰', '；黑夜期间分配'),
    ]
    for anchor, split_at in special:
        n = 0
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    if anchor in c.text:
                        c.text = c.text.split(split_at)[0]
                        n += 1
        print(f'整句删除 [{anchor[:16]}...] -> {n} 处')
        assert n >= 1, '未命中: ' + anchor

    # 表19 R5 C2：修复按钮仅白天可点……保留其余
    n = 0
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if '修复按钮仅白天可点' in c.text:
                    c.text = c.text.split('；修复按钮仅白天可点')[0] + \
                        '；按钮操作的结果立刻反映到弹窗内的状态文字'
                    n += 1
    print(f'修复按钮昼夜限制删除 -> {n} 处')
    assert n >= 1

    d.save(PRD_PATH)
    print('已保存:', PRD_PATH)


if __name__ == '__main__':
    main()
