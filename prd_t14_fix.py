# -*- coding: utf-8 -*-
"""T-14：把代码已实现的 8 大块功能补写进《需求规格说明书 .docx》。

改动 1：4.1 需求范围表新增 8 行
改动 2：4.2 非目标表删除「游戏存档读档」行 + 表后说明
改动 3：6.1 需求总览表新增 FR-08~FR-14（居民疲劳与饥饿 Debuff 合并为 FR-11）
改动 4：新增 6.2.8~6.2.14 详细需求小节（含 11 科技节点表、15 事件表）
改动 5：14.1.1 基础合格验收标准表新增 8 行
一致性修正：10.3 / 12 依赖 / 14.1.2 / 15.2.1 / 16 章中「存档未实现」的残留表述
所有数字均来自源码（出处见改动报告）。
"""
import shutil
import time

import docx

PRD_PATH = r'C:/Users/jeezeqi/Desktop/需求规格说明书 .docx'
BAK_PATH = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/_backup/需求规格说明书_{}.bak.docx'.format(
    time.strftime('%Y%m%d_%H%M%S'))

shutil.copyfile(PRD_PATH, BAK_PATH)
print('已备份:', BAK_PATH)

doc = docx.Document(PRD_PATH)


def find_para(text, startswith=False):
    for p in doc.paragraphs:
        if startswith:
            if p.text.startswith(text):
                return p
        elif text in p.text:
            return p
    raise RuntimeError('未找到段落: ' + text)


# 样式取自现有文档元素
heading_style = find_para('6.2.7', startswith=True).style
fr_dash = find_para('6.2.7', startswith=True).text.split('FR')[1][0]  # 现有 FR 编号中的连字符
table_style = doc.tables[11].style              # 6.2.1 的需求表样式
accept_style = next(p.style for p in doc.paragraphs if p.text.strip() == '验收标准')
given_style = find_para('Given：环境配置正确').style

# ================= 改动 1：4.1 需求范围表（表6）新增 8 行 =================
t41 = doc.tables[6]
rows_41 = [
    ('科技树系统', '11 个科技节点、4 级前置依赖；消耗资源研究；效果含解锁建筑/配方、提升产出；研究「星际航行」触发胜利', 'P0', '代码已实现，文档补写'),
    ('随机事件系统', '每 5~14 个游戏日随机触发 1 个事件；15 种事件、每个含多选项分支；事件弹窗期间暂停每日推进', 'P0', '代码已实现，文档补写'),
    ('建筑升级', '建筑 1~5 级；升级消耗合金与水泥；每级产出 +20%', 'P0', '代码已实现，文档补写'),
    ('居民疲劳', '连续工作累积疲劳（+25/日，上限 100）；疲劳（≥50）时效率与移速 ×0.5；休息恢复（−25/日）', 'P0', '代码已实现，文档补写'),
    ('饥饿 Debuff', '食物归零时全局产出 ×0.5 且每日流失 1 名居民；食物恢复后解除', 'P0', '代码已实现，文档补写'),
    ('猎人工会猎物结算', '猎人工会按交付的猎物数量结算食物；猎物库存每日清零', 'P0', '代码已实现，文档补写'),
    ('资源历史与折线图', '每 10 个游戏日记录一次资源/人口/建筑数快照；可弹窗查看折线图', 'P0', '代码已实现，文档补写'),
    ('存档读档', '保存/读取对局进度到 saves/<槽位>.json；支持列出与删除存档', 'P0', '代码已实现；原 4.2 列为非目标，本次移入范围'),
]
for r in rows_41:
    cells = t41.add_row().cells
    for j, txt in enumerate(r):
        cells[j].text = txt
print('改动1 OK：4.1 表新增 8 行')

# ================= 改动 2：4.2 非目标表（表7）删除「游戏存档读档」行 + 说明 =================
t42 = doc.tables[7]
removed = False
for row in list(t42.rows):
    if '游戏存档读档' in row.cells[0].text:
        t42._tbl.remove(row._tr)
        removed = True
        break
assert removed, '未找到 4.2 存档行'
note42 = doc.add_paragraph('说明：存档读档功能已在 v1.0 代码中实现（见 SaveManager），故自本节移出，'
                           '纳入 4.1 需求范围；受实现限制，存档仅记录资源/人口/已研究科技，不记录建筑地图布局。')
t42._tbl.addnext(note42._p)
print('改动2 OK：4.2 表删除存档行 + 说明')

# ================= 改动 3：6.1 需求总览表（表10）新增 FR-08~FR-14 =================
t61 = doc.tables[10]
rows_61 = [
    ('FR%s08' % fr_dash, '科技树系统', '科技树系统', 'P0', '待测试', '—'),
    ('FR%s09' % fr_dash, '随机事件系统', '随机事件系统', 'P0', '待测试', '—'),
    ('FR%s10' % fr_dash, '建筑升级', '建筑升级', 'P0', '待测试', '—'),
    ('FR%s11' % fr_dash, '居民状态', '居民疲劳与饥饿 Debuff', 'P0', '待测试', '—'),
    ('FR%s12' % fr_dash, '猎人工会', '猎人工会猎物结算', 'P0', '待测试', '—'),
    ('FR%s13' % fr_dash, '数据展示', '资源历史与折线图', 'P0', '待测试', '—'),
    ('FR%s14' % fr_dash, '数据持久化', '存档读档', 'P0', '待测试', '—'),
]
for r in rows_61:
    cells = t61.add_row().cells
    for j, txt in enumerate(r):
        cells[j].text = txt
print('改动3 OK：6.1 表新增 FR-08~FR-14')

# ================= 改动 4：新增 6.2.8~6.2.14 详细需求小节 =================

# 科技树 11 节点（Model/Tech/TechTree.java initializeTechTree()）
TECHS = [
    ('BASIC_MINING', '初级采矿', '解锁矿场', '—', 'FOOD×50', '解锁建筑 QUARRY'),
    ('BASIC_FARMING', '初级农业', '解锁农场', '—', 'WOOD×50', '解锁建筑 FARM'),
    ('BASIC_ENERGY', '初级能源', '解锁能源站', '—', 'STONE×50', '解锁建筑 CEMENTPLANT'),
    ('ADVANCED_MINING', '高级采矿', '采矿效率+50%', 'BASIC_MINING', 'STEEL×100、FOOD×50', 'STEEL 产出 ×1.5'),
    ('METALLURGY', '冶金学', '解锁冶炼厂', 'BASIC_MINING、BASIC_ENERGY', 'STEEL×100、FOOD×50', '解锁建筑 STEELMILL'),
    ('ADVANCED_FARMING', '高级农业', '食物产出+50%', 'BASIC_FARMING', 'STEEL×100、WOOD×50', 'FOOD 产出 ×1.5'),
    ('SUPPLY_CHAIN', '供应链', '解锁建材配方', 'BASIC_MINING、BASIC_FARMING', 'STEEL×100、FOOD×50', '解锁配方 LUMBER_PROCESSING'),
    ('ADVANCED_ENERGY', '高级能源', '能源产出+50%', 'BASIC_ENERGY、METALLURGY', 'STEEL×200、CEMENT×50', 'COAL 产出 ×1.5'),
    ('AI_OPTIMIZATION', '人工智能', '所有产能+20%', 'ADVANCED_MINING、ADVANCED_FARMING', 'STEEL×200、FOOD×100', '全局产出 ×1.2'),
    ('QUANTUM_COMPUTING', '量子计算', '解锁量子实验室', 'ADVANCED_ENERGY、AI_OPTIMIZATION', 'STEEL×500、CEMENT×200', '解锁建筑 TOOLFACTORY'),
    ('INTERSTELLAR_TRAVEL', '星际航行', '胜利条件', 'METALLURGY、ADVANCED_ENERGY、AI_OPTIMIZATION', 'STEEL×1000、CEMENT×500', '触发胜利（TechEffect 为占位实现）'),
]

# 15 种事件（Model/Event/ 各源码文件）
EVENTS = [
    ('👽 外星袭击（AlienAttackEvent）', '建筑数 ≥ 5',
     '⚔️迎战：随机 1 座建筑损坏、1 名居民阵亡；🕊️和谈（消耗 50 食物+30 钢铁）：足额则无损，不足则随机 1 座建筑损坏、2 名居民死亡；🏃撤离：2 名居民失踪'),
    ('🛸 外星商人（AlienTraderEvent）', '第 10 游戏日及以后',
     '🤝交易：50 木材→30 钢铁（木材不足则失败）；🍖交易：20 食物→25 水泥（食物不足则失败）；🚪拒绝：无变化'),
    ('🏛️ 远古遗迹（AncientRuinsEvent）', '建筑数 ≥ 4',
     '🔍考古（消耗 15 工具）：+60 钢铁，工具不足则失败；⛏️掠夺：+50 木材、+50 石头；🚶离开：无变化'),
    ('💰 宇宙赏金（CosmicBountyEvent）', '第 25 游戏日及以后',
     '💵领取：+60 食物、+40 工具；📈投资（消耗 50 钢铁）：+80 水泥，钢铁不足则失败；🙅拒绝：无变化'),
    ('🕊️ 外交任务（DiplomaticMissionEvent）', '人口 ≥ 15',
     '🛸派遣使团（消耗 20 食物+10 工具）：+30 水泥、+30 钢铁，资源不足则失败；🤝婉拒：无变化'),
    ('☄️ 陨石撞击（MeteorStrikeEvent）', '建筑数 ≥ 3',
     '🛡️护盾（消耗 100 钢铁）：足额则建筑无损，不足则目标建筑被摧毁；🚀紧急疏散：目标建筑损坏；💪承受冲击：目标建筑被摧毁'),
    ('👶 人口增长（PopulationGrowthEvent）', '人口 ≥ 5',
     '🍼鼓励生育（消耗 20 食物）：+5 人口，食物不足则搁置；📊保持现状：无变化'),
    ('🏕️ 难民潮（RefugeeWaveEvent）', '第 5 游戏日及以后',
     '🏡接纳（消耗 20 食物）：+5 人口，食物不足则失败；🎁提供物资（消耗 15 食物）：+30 木材、+30 石头，食物不足则失败；🚫拒绝：无变化'),
    ('📦 资源危机（ResourceShortageEvent）', '任一种资源库存为 0',
     '🛒紧急采购（消耗 30 木材）：+15 食物、+15 石头，木材不足则失败；⚫黑市（消耗 5 工具）：+25 食物，工具不足则失败；💪自力更生：无变化'),
    ('☄️ 富饶小行星（RichAsteroidEvent）', '第 5 游戏日及以后',
     '⛏️开采（消耗 10 工具）：+50 铁矿、+30 煤炭，工具不足则失败；🛰️探测器（消耗 20 钢铁）：+40 钢铁、+10 工具，钢铁不足则失败；⏭️忽略：无变化'),
    ('🌞 太阳耀斑（SolarFlareEvent）', '第 15 游戏日及以后且建筑数 ≥ 1',
     '⚡切断电网：全部运行/Boosted 建筑被阻断；🛡️加固（消耗 30 水泥）：足额则无损，不足则随机 1 座建筑损坏；💥承受：随机 1 座建筑损坏'),
    ('🦠 太空瘟疫（SpacePlagueEvent）', '人口 ≥ 10',
     '🔬疫苗（消耗 50 钢铁）：成功则无死亡，不足则 3 名居民死亡；😷隔离（消耗 10 食物）：有效则 2 名居民死亡，食物不足则 5 名居民死亡；🙏听天由命：4 名居民死亡'),
    ('🌪️ 星际风暴（StellarStormEvent）', '第 20 游戏日及以后且建筑数 ≥ 1',
     '🥫储备（消耗 10 食物）：足额则无损，不足则随机 1 座建筑损坏；🧱加固（消耗 30 水泥）：足额则无损，不足则随机 1 座建筑损坏；🎲冒险生产：+40 木材、+30 石头、+20 铁矿，但随机 1 座建筑损坏'),
    ('🔭 科技发现（TechDiscoveryEvent）', '存在可研究科技',
     '🔬深入研究：免费完成一项可研究科技（无则可研究科技时提示无可研究）；📜归档：+20 食物；🗑️丢弃：无变化'),
    ('🚫 贸易禁运（TradeEmbargoEvent）', '第 10 游戏日及以后',
     '🤐走私（消耗 30 食物）：+25 钢铁，食物不足则失败；💬斡旋（消耗 30 木材）：成功则解除禁运，木材不足则禁运持续；😤硬扛：无变化'),
]

SECTIONS = [
    {
        'no': 8, 'fr': 'FR%s08' % fr_dash, 'title': '科技树系统',
        'desc': ('游戏内置科技树，共 11 个科技节点，分 4 级。玩家消耗资源研究科技；'
                 '前置科技已研究且资源充足时方可研究。科技效果包括解锁建筑、解锁配方、提升单资源产出、提升全局产出。'
                 '研究终极科技「星际航行」（INTERSTELLAR_TRAVEL）后触发胜利弹窗。'),
        'cond': '游戏对局正常运行；玩家点击底部「🔬 科技树」按钮打开科技树界面。',
        'rules': ('1. 研究校验：canResearch(id) 同时校验「前置科技已研究」与「资源充足」；research(id) 扣除消耗并应用效果；'
                  'forceResearch(id) 不扣资源（用于事件奖励与读档恢复）。\n'
                  '2. 效果类型共 6 种：UNLOCK_BUILDING（解锁建筑）、BOOST_PRODUCTION（单资源产出加成）、'
                  'REDUCE_CONSUMPTION（降低消耗）、UNLOCK_RECIPE（解锁配方）、INCREASE_CAPACITY（增加容量）、BOOST_ALL（全局产出加成）。\n'
                  '3. 产出倍率 = 单资源倍率 × 全局倍率（getProductionMultiplier）。\n'
                  '4. 胜利条件：研究出 INTERSTELLAR_TRAVEL 即触发胜利弹窗。\n'
                  '5. 已知占位实现：INTERSTELLAR_TRAVEL 的 TechEffect 为占位（unlockBuilding(WOODENCABIN)），'
                  '真正的胜利触发在 GameManager.researchTech 中判定。\n'
                  '6. 11 个节点明细见下表。'),
        'edge': ('1. 资源不足：研究不执行，顶部报错；2. 前置科技未研究：不可研究；'
                 '3. 已研究科技不可重复研究；4. 无任何可研究科技时，事件「科技发现」的深入研究选项提示无可研究科技。'),
        'accept': [
            'Given：对局运行，初级采矿（BASIC_MINING）未研究且食物库存 ≥ 50，When：玩家在科技树界面点击研究「初级采矿」，Then：扣除 50 食物，该科技状态变为已研究，矿场（QUARRY）建筑卡片解除置灰。',
            'Given：对局运行，某科技的前置科技未研究，When：玩家尝试研究该科技，Then：研究不执行，资源数值不变。',
            'Given：对局运行，玩家满足「星际航行」的前置与消耗条件，When：研究「星际航行」成功，Then：弹出胜利弹窗。',
        ],
        'extra_table': ('科技节点明细（11 个）', ['ID', '名称', '说明', '前置', '消耗', '效果'], TECHS),
    },
    {
        'no': 9, 'fr': 'FR%s09' % fr_dash, 'title': '随机事件系统',
        'desc': ('游戏每 5~14 个游戏日随机触发 1 个随机事件；共 15 种事件，每种事件含若干选项分支；'
                 '事件以模态弹窗展示标题、描述与全部选项；事件弹窗打开期间暂停每日推进。'),
        'cond': ('每日结算中调用 eventManager.checkAndTrigger()：turnCounter 递增后，'
                 '当 turnCounter ≥ nextEventTurn 时，从「可触发事件」（canTrigger 通过）中随机选取 1 个触发。'
                 '触发间隔 = 5 + random.nextInt(10) 个游戏日（即 5~14 个游戏日）。'),
        'rules': ('1. 15 种事件的标题、触发条件与各选项后果见下表。\n'
                  '2. 事件弹窗打开期间（eventDialogShowing = true），每日推进短路跳过，避免嵌套事件重入；弹窗关闭后恢复推进。\n'
                  '3. 事件触发时先调用 event.trigger(model) 准备上下文（如随机选择目标建筑），再弹出选项。'),
        'edge': ('1. 无可触发事件时：本轮不弹窗，重置下次触发计时；2. 事件选项执行抛出异常时在弹窗内提示，不导致游戏崩溃。'),
        'accept': [
            'Given：对局已运行 5 个游戏日以上且未打开事件弹窗，When：每日结算执行到事件检测步骤且满足触发间隔，Then：弹出 1 个事件弹窗（标题与描述与下表一致），每日推进暂停。',
            'Given：事件弹窗已打开，When：新的游戏日到来，Then：本轮每日结算跳过，弹窗关闭后恢复推进。',
        ],
        'extra_table': ('15 种随机事件明细', ['事件（类名）', '触发条件', '选项与各选项后果'], EVENTS),
    },
    {
        'no': 10, 'fr': 'FR%s10' % fr_dash, 'title': '建筑升级',
        'desc': ('建筑等级为 1~5 级（MAX_LEVEL = 5）。升级消耗 STEEL = 当前等级×30、CEMENT = 当前等级×15；'
                 '每升 1 级产出 +20%；等级 > 1 时建筑名称后缀「 Lv.N」。'),
        'cond': '游戏对局正常运行；目标建筑等级 < 5；资源满足升级消耗。',
        'rules': ('1. getUpgradeCost()：STEEL = 等级×30，CEMENT = 等级×15；\n'
                  '2. canUpgrade()：等级 < MAX_LEVEL 且资源充足；\n'
                  '3. upgrade()：扣资源、等级 +1；\n'
                  '4. 产出加成（BuildingManager.handle）：levelBonus = 1.0 + 0.2 × (level − 1)，即每升 1 级 +20%。'),
        'edge': ('1. 等级已达 5 级：不可继续升级；2. 资源不足：升级不执行；3. 升级后建筑名称显示「 Lv.N」后缀（N>1）。'),
        'accept': [
            'Given：一座建筑当前等级为 1、产出结算时等级加成系数为 1.0，When：资源充足时执行一次升级，Then：等级变为 2，等级加成系数变为 1.2（+20%），并扣除 STEEL×30、CEMENT×15。',
            'Given：一座建筑当前等级为 5，When：尝试继续升级，Then：升级不执行，资源不变。',
        ],
        'extra_table': None,
    },
    {
        'no': 11, 'fr': 'FR%s11' % fr_dash, 'title': '居民疲劳与饥饿 Debuff',
        'desc': ('居民疲劳：连续工作疲劳度 +25/日（上限 100）；休息 −25/日（下限 0）；疲劳度 ≥ 50 判定为疲劳，'
                 '工作效率 ×0.5、移动速度 ×0.5。饥饿 Debuff：食物库存 ≤ 0 时全局生产效率 ×0.5，且每日流失 1 名居民；'
                 '食物 > 0 后解除。\n'
                 '一致性说明：饥饿对全局产出的影响与原 PRD 6.2.5 文字（"居民死亡只影响人口计数和食物消耗，不影响产出"）'
                 '不一致，以代码为准（代码实际为食物归零时全局产出 ×0.5），本条作为对原设计的调整。'),
        'cond': '游戏日结算执行到 updateFatigue()（全体居民疲劳推进）与 checkFoodAndApplyHunger()（饥饿判定）步骤。',
        'rules': ('1. 疲劳常量（People.java）：FATIGUE_PER_DAY = 25、FATIGUE_RECOVERY_PER_DAY = 25、FATIGUE_MAX = 100、'
                  'FATIGUE_THRESHOLD = 50、FATIGUED_MOVEMENT_SPEED = 0.5、FATIGUED_EFFICIENCY = 0.5；\n'
                  '2. 有工作 → 连续工作天数 +1、疲劳 +25（上限 100）；无工作（休息）→ 连续工作天数清零、疲劳 −25（下限 0）；\n'
                  '3. 疲劳（fatigue ≥ 50）→ getEfficiency() 返回 0.5（工人贡献按效率加权）、移动速度 ×0.5；\n'
                  '4. 饥饿（食物 ≤ 0）→ 全局生产效率倍率 HUNGER_EFFICIENCY = 0.5，每日 killPeople(HUNGER_POPULATION_LOSS_PER_DAY = 1)；\n'
                  '5. 食物 > 0 → 饥饿解除，效率恢复 1.0；\n'
                  '6. 饥饿警告弹窗按「游戏日」节流（hungerWarningDay），每天最多弹一次。'),
        'edge': ('1. 食物长期为 0：每日流失 1 名居民，直至食物恢复或居民全部死亡（居民全部死亡则 gameOver）；'
                 '2. 疲劳与饥饿的效率倍率相乘叠加。'),
        'accept': [
            'Given：一名居民连续工作 2 个游戏日，When：每日结算执行疲劳推进，Then：其疲劳度为 50（2×25），判定为疲劳，工作效率与移动速度均为 ×0.5。',
            'Given：食物库存为 0 且饥饿未激活，When：每日结算执行饥饿判定，Then：全局生产效率倍率变为 0.5，本轮流失 1 名居民，弹出饥饿警告。',
            'Given：饥饿已激活且食物库存恢复为正，When：每日结算执行饥饿判定，Then：饥饿解除，全局生产效率倍率恢复 1.0。',
        ],
        'extra_table': None,
    },
    {
        'no': 12, 'fr': 'FR%s12' % fr_dash, 'title': '猎人工会猎物结算',
        'desc': ('猎人工会（HUNTERGUILD）是 10 种建筑之一，初始即可建造（无需科技解锁），造价为木 10 / 石 10。'
                 '猎人工会自身不自动产出资源；玩家向其交付猎物进入猎物库存，每日结算时按每 1 单位猎物 → 5 单位食物转化，'
                 '并清空猎物库存。'),
        'cond': '地图上已放置猎人工会；玩家通过猎物交付入口（GameManager.deliverPreyToGuild）交付猎物。',
        'rules': ('1. deliverPrey(preyCount)：猎物累加进猎物库存（preyStock）；\n'
                  '2. 每日结算（BuildingManager.handle）：consumePreyStock() 取出并清零库存，'
                  'foodYield = 猎物数 × HUNTER_PREY_FOOD_YIELD（=5）× 饥饿效率倍率；\n'
                  '3. 无猎人工会时交付操作不生效，顶部提示「当前没有猎人工会，无法交付猎物」。'),
        'edge': '猎物库存每日结算后清零，未结算的交付留待下一游戏日结算。',
        'accept': [
            'Given：地图上有一座猎人工会，When：玩家交付 3 单位猎物，Then：下一个游戏日结算后食物 +15（3×5），猎物库存清零。',
            'Given：地图上没有任何猎人工会，When：玩家尝试交付猎物，Then：操作不生效并给出提示。',
        ],
        'extra_table': None,
    },
    {
        'no': 13, 'fr': 'FR%s13' % fr_dash, 'title': '资源历史与折线图',
        'desc': ('每 10 个游戏日记录一次历史快照（day / resources / buildingCount / population），'
                 '保留最近 200 条；玩家点击底部「📊 图表」按钮弹出折线图窗口，展示资源与人口趋势。'),
        'cond': '游戏日结算中 dayCounter % 10 == 0 时执行 recordHistory()；对局运行中点击「📊 图表」按钮。',
        'rules': ('1. 快照字段（HistoricalData.java）：day、resources、buildingCount、population；\n'
                  '2. 历史列表超过 200 条时移除最早一条；\n'
                  '3. ChartView 以折线图弹窗展示（show(Stage)）。'),
        'edge': '游戏运行不足 10 个游戏日时，历史列表可能为空，图表窗口无数据。',
        'accept': [
            'Given：对局运行至第 10 个游戏日结算完成，When：玩家点击「📊 图表」按钮，Then：弹出折线图窗口，且包含第 10 游戏日的资源/人口/建筑数快照。',
        ],
        'extra_table': None,
    },
    {
        'no': 14, 'fr': 'FR%s14' % fr_dash, 'title': '存档读档',
        'desc': ('以 JSON 形式保存/读取对局进度到 saves/<槽位>.json；支持列出存档与删除存档。'
                 '存档字段：day、resources、population、buildingCount、researchedTechs、unlockedRecipes。\n'
                 '已知局限（如实披露）：SaveData 只保存 buildingCount，不保存建筑的种类与地图坐标，读档后建筑布局无法还原。'),
        'cond': '对局运行中点击顶部「💾 保存」/「📂 加载」按钮；读档需存在对应槽位存档文件。',
        'rules': ('1. 存档目录为 saves/，文件名 <slotName>.json（SaveManager.SAVE_DIR = "saves"）；\n'
                  '2. save(model, slotName)：序列化 SaveData 写入 JSON；\n'
                  '3. load(model, slotName)：读入 JSON，恢复天数、资源、人口、已研究科技、已解锁配方（建筑布局无法还原）；\n'
                  '4. listSaves()：列出全部存档；showLoadDialog(model)：弹窗选择存档，支持删除存档。'),
        'edge': ('1. 读档后建筑布局无法还原（存档不记录建筑种类与坐标）；2. 槽位文件损坏或不存在时读取失败并提示。'),
        'accept': [
            'Given：对局运行到第 N 个游戏日，When：玩家点击「💾 保存」，Then：saves/ 目录下生成对应槽位的 JSON 文件。',
            'Given：存在一个有效存档，When：玩家点击「📂 加载」并选择该存档，Then：天数、资源、人口、已研究科技恢复为存档值；建筑布局不还原（已知局限）。',
        ],
        'extra_table': None,
    },
]

# 在 7 业务流程 之前插入（先构建全部元素，再按序搬移）
anchor = find_para('7 业务流程', startswith=True)
built_sections = []
for sec in SECTIONS:
    elems = []
    # 标题
    h = doc.add_paragraph()
    h.style = heading_style
    h.add_run('6.2.%d %s %s' % (sec['no'], sec['fr'], sec['title']))
    elems.append(h)
    # 需求表（4 行 × 2 列）
    tb = doc.add_table(rows=4, cols=2)
    tb.style = table_style
    cells_data = [
        ('需求描述', sec['desc']),
        ('触发条件 / 前置条件', sec['cond']),
        ('功能规则', sec['rules']),
        ('异常与边界', sec['edge']),
    ]
    for i, (k, v) in enumerate(cells_data):
        tb.rows[i].cells[0].text = k
        tb.rows[i].cells[1].text = v
    elems.append(tb)
    # 附表（科技节点 / 事件明细）
    if sec['extra_table']:
        caption, header, rows = sec['extra_table']
        cp = doc.add_paragraph(caption)
        elems.append(cp)
        et = doc.add_table(rows=1, cols=len(header))
        et.style = table_style
        for j, htxt in enumerate(header):
            et.rows[0].cells[j].text = htxt
        for r in rows:
            cells = et.add_row().cells
            for j, txt in enumerate(r):
                cells[j].text = txt
        elems.append(et)
    # 验收标准
    ap = doc.add_paragraph()
    ap.style = accept_style
    ap.add_run('验收标准')
    elems.append(ap)
    for g in sec['accept']:
        gp = doc.add_paragraph(g)
        gp.style = given_style
        elems.append(gp)
    built_sections.append(elems)

# 逆序插入，保证 6.2.8 ~ 6.2.14 顺序正确
for elems in reversed(built_sections):
    for el in elems:
        anchor._p.addprevious(el._p if hasattr(el, '_p') else el._tbl)
print('改动4 OK：新增 6.2.8 ~ 6.2.14 共 %d 个小节' % len(built_sections))

# ================= 改动 5：14.1.1 基础合格验收标准表（表29）新增 8 行 =================
t1411 = doc.tables[29]
rows_1411 = [
    ('科技树系统', '可打开科技树界面；11 个科技节点全部可见；资源充足且前置已研究时可研究；研究「星际航行」后触发胜利弹窗 / 测试方式：手动测试，按步骤复现'),
    ('随机事件系统', '每 5~14 个游戏日触发 1 个事件；15 种事件弹窗展示标题、描述与选项；选项后果与事件描述一致；弹窗期间每日推进暂停 / 测试方式：手动测试，按步骤复现'),
    ('建筑升级', '建筑可升级至 5 级；升级消耗 STEEL=当前等级×30、CEMENT=当前等级×15；每升 1 级产出 +20%；名称显示「 Lv.N」后缀 / 测试方式：手动测试，按步骤复现'),
    ('居民疲劳', '连续工作疲劳 +25/日（上限 100）；疲劳（≥50）时效率与移速 ×0.5；休息 −25/日（下限 0） / 测试方式：手动测试，按步骤复现'),
    ('饥饿 Debuff', '食物归零时全局产出 ×0.5、每日流失 1 名居民；食物恢复后解除，效率恢复 1.0 / 测试方式：手动测试，按步骤复现'),
    ('猎人工会猎物结算', '猎人工会初始可建；交付猎物后每日结算按 1 猎物 → 5 食物转化；猎物库存每日清零 / 测试方式：手动测试，按步骤复现'),
    ('资源历史与折线图', '每 10 个游戏日记录一次快照；图表弹窗展示资源/人口/建筑数折线 / 测试方式：手动测试，按步骤复现'),
    ('存档读档', '可保存/读取/列出/删除 saves/ 下 JSON 存档；读档恢复天数/资源/人口/已研究科技；存档不记录建筑布局（已知局限） / 测试方式：手动测试，按步骤复现'),
]
for r in rows_1411:
    cells = t1411.add_row().cells
    cells[0].text = r[0]
    cells[1].text = r[1]
    cells[2].text = '未开始'
print('改动5 OK：14.1.1 表新增 8 行')

# ================= 一致性修正：清理「存档未实现」残留表述 =================
# 10.3 数据持久化说明
p103 = find_para('10.3')
for p in doc.paragraphs:
    if 'v1.0 版本不实现存档读档功能' in p.text:
        for r in p.runs:
            r.text = ''
        if p.runs:
            p.runs[0].text = ('v1.0 版本已实现存档读档功能（见 6.2.14 FR-14）：对局可保存到 saves/<槽位>.json 并读取恢复；'
                              '存档仅记录资源/人口/建筑数量/已研究科技/已解锁配方，不记录建筑地图布局。')
        print('已修正 10.3 存档表述')
        break
# 12 依赖与约束表（表25）「无存档框架」行
for t in doc.tables:
    for row in t.rows:
        if '无存档框架' in row.cells[0].text:
            row.cells[0].text = '存档框架'
            row.cells[1].text = 'v1.0 已实现 JSON 存档读档（SaveManager）'
            row.cells[2].text = '存档不记录建筑布局，读档后建筑布局无法还原'
            row.cells[3].text = '已实现（建筑布局不还原为已知局限）'
            print('已修正 12 依赖表存档行')
# 15.2.1 暂不支持存档
for p in doc.paragraphs:
    if '暂不支持游戏存档读档功能' in p.text:
        new_text = p.text.replace('；暂不支持游戏存档读档功能', '').replace('暂不支持游戏存档读档功能', '')
        for r in p.runs:
            r.text = ''
        if p.runs:
            p.runs[0].text = new_text
        print('已修正 15.2.1 存档表述')
        break
# 16 章已知未实现清单删除「游戏存档与读档功能。」
for p in doc.paragraphs:
    if p.text.strip() == '游戏存档与读档功能。':
        p._p.getparent().remove(p._p)
        print('已删除 16 章清单中存档条目')
        break
# 14.1.2（表30）删除「数据持久化」行
for t in doc.tables:
    for row in list(t.rows):
        if row.cells[0].text.strip() == '数据持久化':
            t._tbl.remove(row._tr)
            print('已删除 14.1.2 数据持久化行')
            break

doc.save(PRD_PATH)
print('已保存:', PRD_PATH)
