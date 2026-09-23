# -*- coding: utf-8 -*-
"""T-16：清理 PRD 中的不可判定表述（含 T-14 新增内容自查）。

处理靶点（按批次文档）：
- 5.2 核心场景 S-05：「不合理操作」→「非法操作」
- 15.4 游戏核心目标：「合理分配人口工人」→「按建筑容量上限分配居民与工人」
- 15.4「尽可能」句按文档要求保留，并补一句可判定说明
- 13.2 开放问题的「合理」按文档要求不改
"""
import shutil
import time

import docx

PRD_PATH = r'C:/Users/jeezeqi/Desktop/需求规格说明书 .docx'
BAK_PATH = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/_backup/需求规格说明书_{}.bak.docx'.format(
    time.strftime('%Y%m%d_%H%M%S'))

shutil.copyfile(PRD_PATH, BAK_PATH)
print('已备份:', BAK_PATH)

doc = docx.Document(PRD_PATH)

REPLACEMENTS = [
    ('当我做出不合理操作（资源不足建造、超容量分配工人）时',
     '当我执行资源不足建造、超容量分配工人等非法操作时'),
    ('合理分配人口工人', '按建筑容量上限分配居民与工人'),
]


def replace_in_par(p, old, new):
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    p.runs[0].text = p.text.replace(old, new)
    for r in p.runs[1:]:
        r.text = ''
    return True


for old, new in REPLACEMENTS:
    n = 0
    for p in doc.paragraphs:
        if replace_in_par(p, old, new):
            n += 1
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if replace_in_par(p, old, new):
                        n += 1
    print(f'替换 [{old[:20]}...] -> {n} 处')
    assert n >= 1, '未命中: ' + old

# 15.4「尽可能」句保留，并在其后补可判定说明
n = 0
for p in doc.paragraphs:
    if '以尽可能长久维持定居点运转为游玩目标' in p.text:
        p.add_run('以“全部居民死亡且全部建筑停止运转”为结束判定。')
        n += 1
print(f'15.4 补可判定说明 -> {n} 处')
assert n == 1

doc.save(PRD_PATH)
print('已保存:', PRD_PATH)
