# -*- coding: utf-8 -*-
"""T-10（游戏日时长统一）+ T-11（资源种类统一）PRD 修改脚本。

以代码为准修改《需求规格说明书 .docx》：
- T-10: 全文将与游戏日推进时长相关的「1 秒」表述改为 600 现实秒（10 现实分钟）；
  不涉及错误提示展示 3 秒等与游戏日时长无关的表述。
- T-11: 把数据字典 10.1（Resource 资源实体表）重写为代码中实际存在的 9 种资源
  （英文枚举名/中文名/主要来源/主要用途），删除 ore、energy 两条与
  maxValue=500 字段（代码未实现资源上限）。
运行前先备份 PRD 到 docs/_backup/。
"""
import copy
import shutil
import sys

import docx
from docx.oxml.ns import qn

PRD_PATH = r'C:/Users/jeezeqi/Desktop/需求规格说明书 .docx'
BACKUP_PATH = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/_backup/需求规格说明书.v1.0.6.backup.docx'

log = []


def apply_bold_in_run(r, segment):
    """把 run 中 segment 子串拆出并加粗（前后片段继承原 run 属性）。"""
    i = r.text.index(segment)
    before, after = r.text[:i], r.text[i + len(segment):]
    r.text = before
    bold_el = copy.deepcopy(r._r)
    r._r.addnext(bold_el)
    bold_run = docx.text.run.Run(bold_el, r._parent)
    bold_run.text = segment
    bold_run.bold = True
    after_el = copy.deepcopy(r._r)
    bold_el.addnext(after_el)
    after_run = docx.text.run.Run(after_el, r._parent)
    after_run.text = after


def replace_in_par(p, old, new, bold=None):
    """段落/单元格内段落级替换。优先单 run 替换；跨 run 时回退为整段重建。"""
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            if bold and bold in new:
                apply_bold_in_run(r, bold)
            return True
    # 回退：整段重建（格式化信息会丢失，记录日志）
    log.append('整段重建: ' + repr(p.text[:60]))
    p.runs[0].text = p.text.replace(old, new)
    for r in p.runs[1:]:
        r.text = ''
    if bold and bold in new:
        apply_bold_in_run(p.runs[0], bold)
    return True


def iter_all_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def main():
    # 1) 备份
    shutil.copyfile(PRD_PATH, BACKUP_PATH)
    print('已备份:', BACKUP_PATH)

    d = docx.Document(PRD_PATH)

    # ---- T-10: 游戏日时长表述统一（以代码为准：1 游戏日 = 600 现实秒 = 10 现实分钟）----
    replacements = [
        # (旧子串, 新子串, 加粗片段或 None)
        ('游戏一日对应现实 1 秒', '游戏一日对应现实 600 秒（10 现实分钟）', None),              # P190 15.5
        ('当前 1 游戏日 = 1 现实秒', '当前 1 游戏日 = 600 现实秒 = 10 现实分钟', None),         # P236 16
        ('（1 游戏日 = 1 现实秒）', '（1 游戏日 = 600 现实秒 = 10 现实分钟）', None),           # 表6 R1 4.1
        ('1 个游戏日对应现实 1 秒', '1 个游戏日对应现实 600 秒（10 现实分钟）', None),          # 表11 R0 6.2.1
        ('每现实 1 秒，执行一次', '每 600 现实秒（10 现实分钟）推进一个游戏日', '600 现实秒（10 现实分钟）'),  # 表11 R2 6.2.1 功能规则
        ('“游戏一日” 的全部业务逻辑（生产、消耗、人口判定、建筑状态计时）；',
         '。所有资源生产与消耗均按“游戏日”结算，因此现实时间维度上的产出/消耗速率为 1/600，数值平衡无需逐条调整；',
         None),
        ('游戏时间固定为 1 现实秒 = 1 游戏日', '游戏时间固定为 1 游戏日 = 600 现实秒（10 现实分钟）', None),  # 表26 R3 13.1
        ('1 秒 / 日', '600 秒（10 分钟）/ 日', None),                                          # 表27 R2 13.2
        ('不再固定 1 秒一游戏日', '不再固定 600 秒（10 分钟）一游戏日', None),                  # 表30 R4 14.1.2
        ('1 游戏日等价现实 1 秒', '1 游戏日等价 600 现实秒（10 现实分钟）', None),              # 表39 R6 17.2
    ]
    for old, new, bold in replacements:
        n = 0
        for p in iter_all_paragraphs(d):
            if replace_in_par(p, old, new, bold):
                n += 1
        print(f'T-10 替换 [{old[:20]}...] -> 命中 {n} 处')
        assert n >= 1, '未命中: ' + old

    # ---- T-11: 重写数据字典 10.1 资源表（docx 中第 22 个表格，索引 21）----
    t = d.tables[21]
    # 先删掉第 5 列（原「业务说明」列以外的多余列）
    for row in t.rows:
        tr = row._tr
        tr.remove(row.cells[4]._tc)
    # 删除原 3 行数据行（typeName / currentValue / maxValue）
    for _ in range(3):
        t._tbl.remove(t.rows[1]._tr)
    # 表头
    header = ['英文枚举名', '中文名', '主要来源', '主要用途']
    for j, text in enumerate(header):
        t.rows[0].cells[j].text = text
    # 9 种资源（与 Model/Resource/ResourceType.java 逐一对应；来源/用途均取自代码）
    rows = [
        ('FOOD', '食物', '初始库存 25；农场（FARM）每日产出；猎人工会按交付猎物结算'
         '（每 1 单位猎物 → 5 单位食物，见 BuildingManager.HUNTER_PREY_FOOD_YIELD）；'
         'ADVANCED_FOOD 配方转化', '居民每人每日消耗 1 单位；科技研究消耗（如初级采矿消耗 50）'),
        ('WOOD', '木材', '初始库存 100；木屋（WOODENCABIN）每日产出', '建筑建造成本；科技研究消耗'
         '（如初级农业消耗 50）；木材加工厂（LUMBERMILL）每日消耗；LUMBER_PROCESSING 配方输入'),
        ('STONE', '石头', '初始库存 100；采石场（QUARRY）每日产出', '建筑建造成本；科技研究消耗'
         '（如初级能源消耗 50）；水泥厂（CEMENTPLANT）每日消耗；CEMENT_MIXING 配方输入'),
        ('COAL', '煤炭', '初始库存 0；采石场（QUARRY）每日产出；受高级能源科技产出加成（×1.5）',
         '水泥厂（CEMENTPLANT）、冶炼厂（STEELMILL）、工具厂（TOOLFACTORY）每日消耗；'
         'ALLOY_SMELTING、CEMENT_MIXING 配方输入'),
        ('IRON', '铁矿', '初始库存 0；采石场（QUARRY）每日产出', '冶炼厂（STEELMILL）每日消耗；'
         'ALLOY_SMELTING 配方输入'),
        ('STEEL', '合金', '初始库存 0；冶炼厂（STEELMILL）每日产出；ALLOY_SMELTING 配方'
         '（2 铁矿 + 1 煤炭 → 1 合金）', '建筑升级消耗（每级 30×当前等级）；科技研究消耗；'
         '工具厂（TOOLFACTORY）每日消耗'),
        ('CEMENT', '水泥', '初始库存 0；水泥厂（CEMENTPLANT）每日产出；CEMENT_MIXING 配方'
         '（3 石头 + 1 煤炭 → 2 水泥）', '建筑升级消耗（每级 15×当前等级）；科技研究消耗'),
        ('LUMBER', '木材加工品', '初始库存 0；木材加工厂（LUMBERMILL）每日产出；LUMBER_PROCESSING 配方'
         '（2 木材 → 3 木材加工品，由供应链科技解锁）', '当前版本暂无建筑或配方消耗'),
        ('TOOLS', '工具', '初始库存 5；工具厂（TOOLFACTORY）每日产出', '建筑 Boost 强化消耗 1 单位；'
         '建筑修复消耗 1 单位'),
    ]
    for row_data in rows:
        tr = t.add_row()._tr
        if len(tr.tc_lst) > 4:  # tblGrid 仍定义 5 列时，删除多余的单元格
            tr.remove(tr.tc_lst[-1])
        cells = t.rows[-1].cells
        for j, text in enumerate(row_data):
            cells[j].text = text
    # 表后加注：删除 maxValue 字段的说明
    note = d.add_paragraph('注：本表按代码实际实现的 9 种资源重写；原表中的资源存储上限字段'
                           '（maxValue = 500）已删除——当前版本代码未实现资源上限。')
    t._tbl.addnext(note._p)
    print('T-11 资源表已重写为 9 行 + 表头；maxValue 字段已删除并加注说明')

    d.save(PRD_PATH)
    print('已保存:', PRD_PATH)
    for line in log:
        print('  日志:', line)


if __name__ == '__main__':
    main()
