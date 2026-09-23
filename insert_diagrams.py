# -*- coding: utf-8 -*-
"""T-15 插图脚本：向《需求规格说明书 .docx》插入 4 张 PlantUML 图。

- 图 5-1 系统用例图（usecase.png）      -> 5.2 核心场景小节之后（6 功能需求之前）
- 图 7-1 建筑放置流程（activity_place.png）-> 7.1 游戏主流程表格之后
- 图 7-2 每日结算流程（activity_daily.png）-> 7.1.1游戏日结算 之后
- 图 7-3 建筑状态流转（activity_state.png）-> 7.2 建筑状态流转流程 之后
运行前先备份 PRD 到 docs/_backup/。
"""
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

PRD_PATH = r'C:/Users/jeezeqi/Desktop/需求规格说明书 .docx'
BAK_PATH = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/_backup/需求规格说明书.v1.0.6.predocx.bak.docx'
DIAG_DIR = r'C:/Users/jeezeqi/Desktop/FirstVersion/JavaRTS-main/docs/diagrams'


def insert_figure_after(anchor_el, png_path, caption, caption_style=None):
    """在 anchor_el 之后插入 居中图片段落 + 图题段落。"""
    doc_obj = anchor_el.getparent()  # 仅占位，实际使用全局 doc
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(png_path, width=Inches(6))
    p_cap = doc.add_paragraph(caption)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption_style is not None:
        p_cap.style = caption_style
    anchor_el.addnext(p_img._p)
    p_img._p.addnext(p_cap._p)


def insert_figure_before(anchor_par, png_path, caption, caption_style=None):
    """在 anchor_par 之前插入 图片段落 + 图题段落。"""
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(png_path, width=Inches(6))
    p_cap = doc.add_paragraph(caption)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption_style is not None:
        p_cap.style = caption_style
    anchor_par._p.addprevious(p_img._p)
    p_img._p.addnext(p_cap._p)


def find_paragraph(text):
    for p in doc.paragraphs:
        if text in p.text:
            return p
    raise RuntimeError('未找到段落: ' + text)


# 1) 备份
shutil.copyfile(PRD_PATH, BAK_PATH)
print('已备份:', BAK_PATH)

doc = docx.Document(PRD_PATH)

# 图题样式：与现有表题「表 10-1 Resource 资源实体」一致
cap_style = None
for p in doc.paragraphs:
    if 'Resource 资源实体' in p.text:
        cap_style = p.style
        break

# 图 5-1 系统用例图：插在 5.2 核心场景表（第 10 张表，索引 9）之后
insert_figure_after(doc.tables[9]._tbl, DIAG_DIR + '/usecase.png', '图 5-1 系统用例图', cap_style)
print('图 5-1 已插入（5.2 核心场景表之后）')

# 图 7-1 建筑放置流程：插在 7.1 游戏主流程表（第 19 张表，索引 18）之后
insert_figure_after(doc.tables[18]._tbl, DIAG_DIR + '/activity_place.png', '图 7-1 建筑放置流程', cap_style)
print('图 7-1 已插入（7.1 游戏主流程表之后）')

# 图 7-2 每日结算流程：插在 7.1.1 小节末段之后
p_last_711 = find_paragraph('全局食物消耗与居民死亡判定')
insert_figure_after(p_last_711._p, DIAG_DIR + '/activity_daily.png', '图 7-2 每日结算流程', cap_style)
print('图 7-2 已插入（7.1.1游戏日结算之后）')

# 图 7-3 建筑状态流转：插在 7.2 小节之后（8 界面与交互需求 之前）
p_ch8 = find_paragraph('8 界面与交互需求')
insert_figure_before(p_ch8, DIAG_DIR + '/activity_state.png', '图 7-3 建筑状态流转', cap_style)
print('图 7-3 已插入（7.2 建筑状态流转流程之后）')

doc.save(PRD_PATH)
print('已保存:', PRD_PATH)
